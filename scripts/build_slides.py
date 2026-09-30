"""Sinh slide thuyết trình (~12 trang) từ reports/results.json + figures.

Chạy sau notebook: python -m scripts.build_slides
"""
from __future__ import annotations

import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.config import FIGURES_DIR, LABELS, ROOT

RESULTS = ROOT / "reports" / "results.json"
AGREEMENT = ROOT / "reports" / "agreement.json"
OUT = ROOT / "reports" / "Slide_BaoCao_DoAn.pptx"
NAVY = RGBColor(0x1A, 0x36, 0x5D)
ACCENT = RGBColor(0xDD, 0x6B, 0x20)
GREY = RGBColor(0x4A, 0x55, 0x68)
MODEL_VI = {"MultinomialNB": "Naive Bayes", "LogisticRegression": "Logistic Reg.", "LinearSVC": "LinearSVC",
            "RandomForest": "Random Forest"}


def pct(x) -> str:
    return "–" if x is None else f"{100 * float(x):.1f}%"


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.blank = self.prs.slide_layouts[6]

    def slide(self, title: str, subtitle: str | None = None):
        s = self.prs.slides.add_slide(self.blank)
        bar = s.shapes.add_shape(1, 0, 0, self.prs.slide_width, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.9)).text_frame
        tb.text = title
        r = tb.paragraphs[0].runs[0]; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = NAVY
        if subtitle:
            p = tb.add_paragraph(); p.text = subtitle
            p.runs[0].font.size = Pt(15); p.runs[0].font.color.rgb = GREY
        return s

    def closing_slide(self, title: str, subtitle: str):
        """Slide kết: tiêu đề + phụ đề căn giữa cả ngang lẫn dọc."""
        s = self.prs.slides.add_slide(self.blank)
        bar = s.shapes.add_shape(1, 0, 0, self.prs.slide_width, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
        tf = s.shapes.add_textbox(0, 0, self.prs.slide_width, self.prs.slide_height).text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, (text, size, bold, color) in enumerate([(title, 36, True, NAVY), (subtitle, 22, False, GREY)]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = text
            p.alignment = PP_ALIGN.CENTER
            if i: p.space_before = Pt(12)
            r = p.runs[0]; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        return s

    def title_slide(self, title: str, subtitle: str, advisor: str, group_num: str,
                    members: list[tuple[str, str]], highlights: list[str]):
        s = self.prs.slides.add_slide(self.blank)

        # 1. Dải màu cam Accent ở đỉnh slide
        bar = s.shapes.add_shape(1, 0, 0, self.prs.slide_width, Inches(0.14))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

        # 2. Header trường & môn học
        tb_hdr = s.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.733), Inches(0.7)).text_frame
        tb_hdr.word_wrap = True
        p0 = tb_hdr.paragraphs[0]
        p0.text = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN – ĐHQG-HCM"
        p0.runs[0].font.size = Pt(12.5)
        p0.runs[0].font.bold = True
        p0.runs[0].font.color.rgb = NAVY
        p1 = tb_hdr.add_paragraph()
        p1.text = "BÁO CÁO ĐỒ ÁN MÔN HỌC: MÁY HỌC (CS114.F31.CN2.TTNT)"
        p1.runs[0].font.size = Pt(11.5)
        p1.runs[0].font.bold = True
        p1.runs[0].font.color.rgb = ACCENT
        p1.space_before = Pt(2)

        # 3. Khối tiêu đề đề tài chính
        tb_title = s.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.733), Inches(1.65)).text_frame
        tb_title.word_wrap = True
        pt = tb_title.paragraphs[0]
        pt.text = title
        pt.runs[0].font.size = Pt(30)
        pt.runs[0].font.bold = True
        pt.runs[0].font.color.rgb = NAVY
        psub = tb_title.add_paragraph()
        psub.text = subtitle
        psub.runs[0].font.size = Pt(15.5)
        psub.runs[0].font.color.rgb = GREY
        psub.space_before = Pt(6)

        # 4. Hai thẻ thông tin cân đối ở nửa dưới slide
        c_top, c_w, c_h = 3.05, 5.68, 3.85
        card_bg = RGBColor(0xF8, 0xFA, 0xFC)
        card_border = RGBColor(0xCB, 0xD5, 0xE1)

        # Thẻ bên trái: Thông tin thực hiện & Thành viên
        c1_left = 0.8
        card1 = s.shapes.add_shape(5, Inches(c1_left), Inches(c_top), Inches(c_w), Inches(c_h))
        card1.fill.solid(); card1.fill.fore_color.rgb = card_bg
        card1.line.color.rgb = card_border; card1.line.width = Pt(1)

        tf1 = s.shapes.add_textbox(Inches(c1_left + 0.35), Inches(c_top + 0.28), Inches(c_w - 0.7), Inches(c_h - 0.55)).text_frame
        tf1.word_wrap = True
        h1 = tf1.paragraphs[0]
        h1.text = "THÔNG TIN THỰC HIỆN"
        h1.runs[0].font.size = Pt(15); h1.runs[0].font.bold = True; h1.runs[0].font.color.rgb = NAVY
        h1.space_after = Pt(12)

        p_adv_lbl = tf1.add_paragraph()
        p_adv_lbl.text = "Giảng viên hướng dẫn:"
        p_adv_lbl.runs[0].font.size = Pt(12); p_adv_lbl.runs[0].font.color.rgb = GREY
        p_adv = tf1.add_paragraph()
        p_adv.text = advisor
        p_adv.runs[0].font.size = Pt(14); p_adv.runs[0].font.bold = True; p_adv.runs[0].font.color.rgb = NAVY
        p_adv.space_after = Pt(10)

        p_grp = tf1.add_paragraph()
        p_grp.text = f"Nhóm sinh viên thực hiện: {group_num}"
        p_grp.runs[0].font.size = Pt(13); p_grp.runs[0].font.bold = True; p_grp.runs[0].font.color.rgb = ACCENT
        p_grp.space_after = Pt(6)

        for name, mssv in members:
            pm = tf1.add_paragraph()
            pm.text = f"•  {name}  –  {mssv}"
            pm.runs[0].font.size = Pt(13.5); pm.runs[0].font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
            pm.space_after = Pt(4)

        # Thẻ bên phải: Tổng quan giải pháp & Điểm nổi bật
        c2_left = 6.85
        card2 = s.shapes.add_shape(5, Inches(c2_left), Inches(c_top), Inches(c_w), Inches(c_h))
        card2.fill.solid(); card2.fill.fore_color.rgb = card_bg
        card2.line.color.rgb = card_border; card2.line.width = Pt(1)

        tf2 = s.shapes.add_textbox(Inches(c2_left + 0.35), Inches(c_top + 0.28), Inches(c_w - 0.7), Inches(c_h - 0.55)).text_frame
        tf2.word_wrap = True
        h2 = tf2.paragraphs[0]
        h2.text = "TỔNG QUAN GIẢI PHÁP"
        h2.runs[0].font.size = Pt(15); h2.runs[0].font.bold = True; h2.runs[0].font.color.rgb = NAVY
        h2.space_after = Pt(12)

        for it in highlights:
            ph = tf2.add_paragraph()
            ph.text = f"•  {it}"
            ph.runs[0].font.size = Pt(13); ph.runs[0].font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
            ph.space_after = Pt(8)

        return s

    def bullets(self, s, items, left=0.7, top=1.6, width=12, size=19):
        tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.5)).text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            sub = it.startswith("  ")
            p.text = ("– " if sub else "• ") + it.strip()
            p.level = 1 if sub else 0
            p.space_after = Pt(8)
            for r in p.runs:
                r.font.size = Pt(size - 3 if sub else size); r.font.color.rgb = GREY if sub else RGBColor(0x1A, 0x20, 0x2C)

    def image(self, s, name, left, top, width=None, height=None):
        path = FIGURES_DIR / f"{name}.png"
        if path.exists():
            s.shapes.add_picture(str(path), Inches(left), Inches(top),
                                 width=Inches(width) if width else None, height=Inches(height) if height else None)

    def table(self, s, header, rows, left, top, width, row_h=0.42, size=13):
        t = s.shapes.add_table(len(rows) + 1, len(header), Inches(left), Inches(top), Inches(width),
                               Inches(row_h * (len(rows) + 1))).table
        for j, h in enumerate(header):
            c = t.cell(0, j); c.text = str(h)
            c.text_frame.paragraphs[0].runs[0].font.size = Pt(size); c.text_frame.paragraphs[0].runs[0].font.bold = True
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                c = t.cell(i, j); c.text = str(v)
                c.text_frame.paragraphs[0].runs[0].font.size = Pt(size)

    def save(self):
        self.prs.save(OUT)


def main() -> None:
    R = json.loads(RESULTS.read_text(encoding="utf-8"))
    M, best = R["models"], R["best_model"]
    has_real = R.get("n_real", 0) > 0
    n_human = sum(s == "human" for s in json.loads(AGREEMENT.read_text(encoding="utf-8")).get("sources", {}).values()) \
        if AGREEMENT.exists() else 0
    d = Deck()

    # Slide 1: Trang bìa chuyên nghiệp
    d.title_slide(
        title="Tra cứu Luật An toàn, vệ sinh lao động bằng NLP tiếng Việt",
        subtitle="Hệ thống hai giai đoạn: Phân loại 8 nhóm quy định (ML) & Tra cứu Điều luật cụ thể (IR)",
        advisor="ThS. Cáp Phạm Đình Thăng",
        group_num="Nhóm 27",
        members=[
            ("Đỗ Quốc Hoàng", "26410043"),
            ("Nguyễn Trí Toàn", "26410135"),
            ("Nguyễn Văn Thái", "26410108"),
        ],
        highlights=[
            "Bài toán: Câu hỏi đời thường → 8 nhóm quy định (ML) → Top-k Điều luật (IR)",
            "Văn bản pháp lý: Luật 84/2015/QH13, VBHN 14/2024 (phạm vi Điều 1–62)",
            f"Mô hình chọn theo val: TF-IDF (1,2)-gram + {best} (Test F1: {100 * M[best]['test_macro_f1']:.1f}%"
            + (f", Thật F1: {100 * M[best]['real_macro_f1']:.1f}%)" if has_real else ")"),
            "Cam kết kỹ thuật: 100% Machine Learning truyền thống, chạy offline & demo Streamlit",
        ],
    )

    s = d.slide("Bối cảnh & kiến trúc hai giai đoạn")
    d.bullets(s, ["Người lao động hỏi bằng ngôn ngữ đời thường, không nhớ số Điều, Khoản",
                  "Tìm kiếm từ khóa chính xác: dễ không ra kết quả hoặc ra quá nhiều",
                  "Giai đoạn 1 (ML): TF-IDF + bộ phân loại → 1 trong 8 nhóm quy định",
                  "Giai đoạn 2 (IR): cosine similarity giữa câu hỏi và các Điều trong nhóm → Top-k Điều luật",
                  "Không dùng LLM/API bên ngoài để trả lời: toàn bộ là mô hình học máy truyền thống"])

    s = d.slide("Hệ thống 8 lớp nhãn", "Mỗi Điều thuộc đúng 1 lớp; nhãn câu hỏi suy ra từ Điều trả lời trực tiếp")
    d.table(s, ["Lớp", "Mã", "Nội dung", "Số Điều"],
            [[f"C{k}", v[0], v[1], len(v[2])] for k, v in LABELS.items()], 0.6, 1.7, 12.1, size=13)

    s = d.slide("Dữ liệu", f"{R['n_samples']} câu hỏi · {R['n_seeds']} câu gốc × 5 văn phong"
                + (f" · {R['n_real']} câu hỏi thật" if has_real else ""))
    sz = R["split_sizes"]
    d.bullets(s, [f"Synthetic do LLM (Claude) sinh từ nội dung Điều luật: 40 câu gốc/lớp, 5 văn phong; nhãn suy từ Điều, kiểm bằng validator",
                  "Cân bằng lớp: 200 mẫu mỗi lớp",
                  f"Chia theo nhóm câu gốc (StratifiedGroupKFold): train {sz.get('train')} · val {sz.get('val')} · test {sz.get('test')}",
                  "  Mọi paraphrase của một câu gốc nằm cùng 1 tập → không rò rỉ",
                  ] + ([f"Test thật: {R['n_real']} câu của người dân trên chinhsachonline.chinhphu.vn, nhãn LLM gán theo Điều mà Bộ trả lời viện dẫn"
                        + (f", {n_human} thành viên gán nhãn mù kiểm chứng" if n_human else "")] if has_real else []),
              width=6.3, size=17)
    d.image(s, "eda_class_distribution", 7.0, 1.7, width=6.0)

    s = d.slide("EDA & tiền xử lý")
    d.image(s, "eda_wordcloud", 0.4, 1.4, width=7.6)
    d.bullets(s, ["NFC, chữ thường, bỏ URL/ký tự đặc biệt",
                  "Tách từ underthesea: bảo hộ lao động → bảo_hộ_lao_động",
                  "Lọc stopwords, giữ từ phủ định (không, chưa, cấm, được)",
                  "TF-IDF (1,2)-gram, sublinear_tf, min_df=2",
                  "Vectorizer nằm trong Pipeline: fit chỉ trên train"], left=8.2, width=4.9, size=16)

    s = d.slide("So sánh 4 mô hình", "GridSearchCV + StratifiedGroupKFold(5), tối ưu Macro-F1; chọn theo validation")
    hdr = ["Mô hình", "CV F1", "Train F1", "Val F1", "Test F1"] + (["Thật F1"] if has_real else [])
    d.table(s, hdr, [[MODEL_VI[m], pct(x["cv_macro_f1"]), pct(x["train_macro_f1"]), pct(x["val_macro_f1"]),
                      pct(x["test_macro_f1"])] + ([pct(x.get("real_macro_f1"))] if has_real else []) for m, x in M.items()],
            0.6, 1.8, 7.0, size=14)
    d.image(s, "model_comparison", 7.8, 1.8, width=5.3)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(4.4), Inches(7), Inches(1)).text_frame
    tb.text = f"Tốt nhất: {MODEL_VI[best]} (val Macro-F1 {pct(M[best]['val_macro_f1'])})"
    tb.paragraphs[0].runs[0].font.size = Pt(22); tb.paragraphs[0].runs[0].font.bold = True
    tb.paragraphs[0].runs[0].font.color.rgb = ACCENT

    s = d.slide("Ma trận nhầm lẫn", f"{MODEL_VI[best]} · test synthetic" + (" và câu hỏi thật" if has_real else ""))
    d.image(s, "confusion_matrix", 0.5, 1.5, width=12.3)

    s = d.slide("Ablation study", f"{MODEL_VI[best]} · mỗi cặp chỉ khác 1 yếu tố")
    ab = R["ablation"]
    d.table(s, ["Cấu hình", "CV F1 ± std", "Test F1"] + (["Thật F1"] if has_real else []),
            [[k, f"{pct(v['cv_f1_mean'])} ± {100 * v['cv_f1_std']:.1f}", pct(v["test_f1"])]
             + ([pct(v.get("real_f1"))] if has_real else []) for k, v in ab.items()], 0.5, 1.7, 8.2, size=13)
    d.image(s, "ablation", 8.9, 1.7, width=4.2)

    s = d.slide("Phân tích lỗi")
    tc = list(R["top_confusions"].items())[:4]
    d.bullets(s, ["Cặp hay nhầm nhất (4 mô hình, test):"] + [f"  {k}: {v} lần" for k, v in tc]
              + ["Accuracy theo văn phong: " + ", ".join(f"{k} {pct(v)}" for k, v in R["acc_by_style"].items())]
              + ([f"Câu hỏi thật dài ~{R.get('real_len_mean', 0):.0f} âm tiết, nhiều bối cảnh và dẫn chiếu Nghị định → domain shift"] if has_real else []),
              width=7.2, size=17)
    d.image(s, "acc_by_style", 7.9, 1.8, width=5.1)

    s = d.slide("Tra cứu Điều luật (giai đoạn 2)")
    ret = R["retrieval"]
    cols = list(ret)
    col_vi = {"test": "Test synthetic", "real": "Câu hỏi thật"}
    row_vi = {"oracle_top1": "Oracle Top-1", "e2e_top1": "End-to-end Top-1", "nofilter_top1": "No-filter Top-1",
              "oracle_top3": "Oracle Top-3", "e2e_top3": "End-to-end Top-3", "nofilter_top3": "No-filter Top-3",
              "oracle_top1_multi_article_classes": "Oracle Top-1 (lớp nhiều Điều)"}
    d.table(s, ["Chỉ số"] + [col_vi.get(c, c) for c in cols],
            [[row_vi.get(k, k)] + [pct(ret[c][k]) for c in cols] for k in ret[cols[0]]], 0.6, 1.6, 7.5, size=13)
    d.bullets(s, ["Oracle: biết đúng lớp → đo riêng retrieval", "End-to-end: dùng lớp dự đoán",
                  "No-filter: tìm trên cả 62 Điều (baseline)", "Demo: streamlit run app.py"], left=8.4, width=4.6, size=16)

    s = d.slide("Kết luận & hướng phát triển")
    d.bullets(s, [f"Pipeline hoàn chỉnh: corpus 93 Điều (2024) → {R['n_samples']} câu hỏi → 4 mô hình → ablation → tra cứu → demo",
                  f"{MODEL_VI[best]} tốt nhất; chia theo nhóm câu gốc để số liệu không bị thổi phồng",
                  "Hạn chế: dữ liệu huấn luyện synthetic, test thật nhỏ, chưa xử lý câu nhiều ý / ngoài phạm vi, chưa có mức phạt",
                  "Hướng phát triển: thêm dữ liệu thật + kappa, multi-label, Đ63–93 và Nghị định, NĐ 12/2022 mức phạt, PhoBERT"])

    d.closing_slide("Cảm ơn thầy và các bạn đã lắng nghe", "Hỏi & đáp")
    d.save()
    print(f"-> {OUT} ({len(d.prs.slides)} slides)")


if __name__ == "__main__":
    main()
